
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

for fname in ['Researcher_2_Day_1_Domain_Matrix_REVIEWED.xlsx', 'Researcher_2_Day_2_Laycan_Matrix.xlsx']:
    wb = openpyxl.load_workbook(fname)
    print('='*70)
    print('WORKBOOK:', fname)
    print('Sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        ws = wb[sh]
        print(f'\n--- Sheet: {sh} ---')
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                print(row)
